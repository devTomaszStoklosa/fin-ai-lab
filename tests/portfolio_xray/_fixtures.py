from io import BytesIO

import openpyxl

XTB_OPEN_POSITIONS_HEADERS = [
    "Product",
    "Instrument/Position",
    "Ticker",
    "Category",
    "Type",
    "Volume",
    "Value",
    "Current price",
    "Open price",
    "Open time (UTC)",
    "Stop Loss",
    "Take Profit",
    "Net Profit %",
    "Net Profit",
    "Gross Profit",
    "Margin",
    "Open Commission",
    "Swap",
    "Rollover",
]


def build_synthetic_xtb_workbook(extra_note: str | None = None) -> bytes:
    """A small workbook shaped like the real XTB export: metadata rows before
    the header, then summary rows (Category set) interleaved with per-lot
    detail rows (Type set, Category blank) for the same instrument.

    extra_note, when given, is appended as a trailing cell on the first
    summary row — past the declared headers, so it never becomes a mapped
    field, only something a whole-file scan (e.g. injection flagging) sees.
    """
    workbook = openpyxl.Workbook()
    workbook.active.title = "Closed Positions"
    workbook.create_sheet("Cash Operations")
    sheet = workbook.create_sheet("Open Positions")

    for _ in range(8):
        sheet.append([])

    sheet.append(XTB_OPEN_POSITIONS_HEADERS)
    padding = [None] * (len(XTB_OPEN_POSITIONS_HEADERS) - 9)
    first_row = [
        "My Trades", "MSCI ACWI", "ISAC.UK", "ETF", None, 6.0, 2745.41, None, 69.48
    ] + padding
    if extra_note is not None:
        first_row.append(extra_note)
    sheet.append(first_row)
    sheet.append(
        ["My Trades", "1140380016", "ISAC.UK", None, "BUY", 6.0, 2745.41, 122.24, 69.48] + padding
    )
    sheet.append(
        ["My Trades", "Atrem", "ATR.PL", "STC", None, 1.0, 53.3, None, 12.25] + padding
    )
    sheet.append(
        ["My Trades", "987654321", "ATR.PL", None, "BUY", 1.0, 53.3, 53.3, 12.25] + padding
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# Valid ISIN checksums for a made-up "SYNTH" issuer — not a real company,
# picked only to satisfy canonical.Position's checksum validation.
SYNTH_A_ISIN = "PLSYNTH00016"
SYNTH_B_ISIN = "PLSYNTH00024"


def build_synthetic_bossa_csv() -> bytes:
    """Shaped like the real Bossa "historia operacji" export (hisPW.csv):
    semicolon-delimited, cp1250, one row per transaction, "-" column carries
    K (buy) / S (sell). SYNTHA is bought then partly sold (stays open,
    weighted-average cost updates); SYNTHB is bought then fully sold (closes
    to zero, must not appear in the resulting positions)."""
    rows = [
        "data;papier;isin;ilość;-;cena;wartość;prowizja;po prowizji;waluta",
        f"01.01.2026 10:00:00;SYNTHA;{SYNTH_A_ISIN};10;K;100,00;1000,00;5,00;1005,00;PLN",
        f"05.01.2026 10:00:00;SYNTHA;{SYNTH_A_ISIN};4;S;110,00;440,00;2,00;438,00;PLN",
        f"01.01.2026 11:00:00;SYNTHB;{SYNTH_B_ISIN};5;K;50,00;250,00;1,00;251,00;PLN",
        f"02.01.2026 11:00:00;SYNTHB;{SYNTH_B_ISIN};5;S;60,00;300,00;1,00;299,00;PLN",
    ]
    return ("\r\n".join(rows) + "\r\n").encode("cp1250")
